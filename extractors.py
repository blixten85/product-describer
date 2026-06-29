"""Turns an uploaded file into a list of product rows.

CSV/Excel already have structured columns (Site, Product, Price (SEK)) and
are parsed directly. Free-form formats (txt, docx, pdf) don't, so their text
is sent to the AI provider chain once per file with an extraction prompt
that asks it to find every item mentioned and return them as JSON.
"""

import csv
import json
import logging
import os
import re
from pathlib import Path

from providers import AllProvidersExhausted, ProviderChain

_log = logging.getLogger("describer.extractors")

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".txt", ".docx", ".pdf"}

# Tak mot resursutmattning via en illvillig/jättestor uppladdning. PDF-sidor
# läses tungt av pdfplumber; AI-extraktionen kapar ändå texten till MAX_EXTRACT_CHARS.
MAX_PDF_PAGES = int(os.getenv("MAX_PDF_PAGES", "200"))
MAX_EXTRACT_CHARS = int(os.getenv("MAX_EXTRACT_CHARS", "50000"))

ROW_FIELDS = ["Site", "Product", "Price (SEK)", "Link"]

_JSON_ARRAY = re.compile(r"\[.*\]", re.DOTALL)

EXTRACTION_PROMPT = (
    "Du får ett textdokument. Hitta varje enskild produkt/pryl som nämns i texten. "
    "Svara ALLTID med endast en giltig JSON-array, utan kodstaket eller extra text, "
    "i exakt detta format:\n"
    '[{"Product": "...", "Site": "...", "Price (SEK)": "..."}]\n'
    "- 'Product' (krävs): produktens namn.\n"
    "- 'Site' och 'Price (SEK)' (valfria): lämna som tom sträng om okänt.\n"
    "Hitta om möjligt ALLA produkter i dokumentet, inte bara de första."
)


class ExtractionError(Exception):
    pass


def extract_rows(path: str, chain: ProviderChain | None = None) -> tuple[list[dict], list[str]]:
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix == ".xlsx":
        return _parse_excel(path)
    if suffix in (".txt", ".docx", ".pdf"):
        if chain is None:
            raise ExtractionError(
                f"{suffix}-filer kräver en konfigurerad AI-leverantör för att hitta produkter."
            )
        text = _extract_text(path, suffix)
        return _ai_extract(text, chain)
    raise ExtractionError(f"Filtypen {suffix} stöds inte.")


def _parse_csv(path: str) -> tuple[list[dict], list[str]]:
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        fieldnames = list(reader.fieldnames or [])
    return rows, fieldnames


def _parse_excel(path: str) -> tuple[list[dict], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = [str(c) if c is not None else "" for c in next(rows_iter)]
    except StopIteration:
        return [], ROW_FIELDS
    rows = []
    for raw_row in rows_iter:
        if all(c is None for c in raw_row):
            continue
        rows.append({header[i]: ("" if v is None else str(v)) for i, v in enumerate(raw_row) if i < len(header)})
    return rows, header


def _extract_text(path: str, suffix: str) -> str:
    if suffix == ".txt":
        return Path(path).read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        import docx

        document = docx.Document(path)
        return "\n".join(p.text for p in document.paragraphs)
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > MAX_PDF_PAGES:
                _log.warning("PDF har %d sidor, läser bara de första %d", len(pdf.pages), MAX_PDF_PAGES)
            return "\n".join(page.extract_text() or "" for page in pdf.pages[:MAX_PDF_PAGES])
    raise ExtractionError(f"Filtypen {suffix} stöds inte.")


def _ai_extract(text: str, chain: ProviderChain) -> tuple[list[dict], list[str]]:
    text = text.strip()
    if not text:
        raise ExtractionError("Dokumentet innehöll ingen text.")

    if len(text) > MAX_EXTRACT_CHARS:
        _log.warning(
            "dokumentet kapades till %d av %d tecken inför AI-extraktion — produkter därefter kan saknas",
            MAX_EXTRACT_CHARS, len(text),
        )
    content = chain.call(EXTRACTION_PROMPT, text[:MAX_EXTRACT_CHARS])

    match = _JSON_ARRAY.search(content or "")
    if not match:
        raise ExtractionError("AI-leverantören kunde inte hitta några produkter i dokumentet.")
    try:
        items = json.loads(match.group(0))
    except json.JSONDecodeError as e:
        raise ExtractionError("AI-leverantörens svar gick inte att tolka som JSON.") from e

    rows = [
        {
            "Site": str(item.get("Site", "")).strip(),
            "Product": str(item.get("Product", "")).strip(),
            "Price (SEK)": str(item.get("Price (SEK)", "")).strip(),
            "Link": "",
        }
        for item in items
        if str(item.get("Product", "")).strip()
    ]
    if not rows:
        raise ExtractionError("Inga produkter kunde identifieras i dokumentet.")
    return rows, ROW_FIELDS
